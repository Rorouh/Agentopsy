"""Exportación del timeline de investigación como hoja de cálculo.

Reusa el builder existente (``assemble_investigation_timeline``) para producir los
eventos y verifica dos cosas distintas:

- que la hoja se ABRE bien donde el perito la abre: un `.xlsx` real, con acentos
  intactos y las celdas ya separadas, sin depender del idioma del sistema de quien
  lo abra (el CSV que había antes dependía de las dos cosas y perdía una: con la
  declaración ``sep=;`` Excel deja de aplicar el BOM y el texto sale ilegible), y
  un bloque de procedencia que dice de qué caso es la tabla, cuándo se exportó y
  qué trae;
- que los datos siguen siendo honestos: una fila por evento de los dos tipos, el
  argv literal auditado, y un caso sin actividad da la cabecera sin filas.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from forensia.export_hoja import NO_APLICA
from forensia.findings.store import Finding
from forensia.timeline.builder import assemble_investigation_timeline
from forensia.timeline.export import HOJA_HEADER, timeline_to_hoja

#: El bloque de procedencia de la hoja del timeline (lo que hay antes de la tabla).
_PROCEDENCIA_FILAS = 10

_IDX = {name: i for i, name in enumerate(HOJA_HEADER)}


def _hoja(blob: bytes) -> list[list[Any]]:
    """Las celdas tal cual las lee una hoja de cálculo, conservando el tipo."""
    ws = load_workbook(io.BytesIO(blob)).active
    return [[c.value for c in fila] for fila in ws.iter_rows()]


def _rows(blob: bytes) -> list[list[str]]:
    """La TABLA sola, en texto, para comparar contra el contrato de la cabecera."""
    tabla = _hoja(blob)[_PROCEDENCIA_FILAS + 1 :]
    return [
        ["" if v is None else str(v) for v in fila]
        for fila in tabla
        if any(v is not None for v in fila)
    ]


def _finding(fid: str, created_at: str) -> Finding:
    return Finding(
        id=fid,
        case_id="c1",
        title="Persistencia detectada",
        summary="run key sospechosa",
        severity="high",
        evidence_id="e1",
        tool_id="RegRipper",
        run_id=None,
        created_at=created_at,
        mitre_hints=["T1547.001"],
    )


def _audit_run() -> list[dict[str, Any]]:
    return [
        {
            "action": "tool_run_start",
            "ts_utc": "2026-07-15T10:00:00+00:00",
            "run_id": "r1",
            "tool_id": "tsk_fls",
            "argv": ["fls", "-m", "/", "img.raw"],
            "evidence_id": "e1",
        },
        {
            "action": "tool_run_finish",
            "ts_utc": "2026-07-15T10:00:05+00:00",
            "run_id": "r1",
            "status": "finished",
            "exit_code": 0,
            "output_files_count": 0,
        },
    ]


def test_empty_timeline_is_header_only() -> None:
    rows = _rows(timeline_to_hoja([]))
    assert rows == [list(HOJA_HEADER)]


def test_sheet_opens_as_a_spreadsheet_and_declares_its_provenance() -> None:
    events = assemble_investigation_timeline(
        _audit_run(), [_finding("f1", "2026-07-15T11:00:00.000Z")]
    )
    blob = timeline_to_hoja(
        events,
        case_id="c1",
        case_name="Caso Ñandú",
        exported_at="2026-08-06T13:05:42Z",
    )
    # Un `.xlsx` es un paquete OOXML: empieza por la firma de un zip. Es lo que
    # hace que el sistema lo abra con la hoja de cálculo y no con un editor.
    assert blob[:2] == b"PK"
    hoja = _hoja(blob)
    procedencia = dict((r[0], r[1]) for r in hoja[:_PROCEDENCIA_FILAS])
    assert procedencia["Caso"] == "Caso Ñandú"  # UTF-8 de verdad, no transliterado
    assert procedencia["Identificador del caso"] == "c1"
    assert procedencia["Exportado (UTC)"] == "2026-08-06T13:05:42Z"
    assert procedencia["Eventos en la hoja"] == "2"
    assert "1 ejecuciones de herramienta (0 con error)" in procedencia["Composición"]
    assert procedencia["Primer evento"].startswith("2026-07-15T10:00:00")
    assert procedencia["Último evento"].startswith("2026-07-15T11:00:00")
    # La línea vacía es el contrato para quien lea la hoja con un programa.
    assert all(v is None for v in hoja[_PROCEDENCIA_FILAS])


def test_sheet_has_one_row_per_event_across_both_kinds() -> None:
    audit = _audit_run()
    audit[1] |= {"stdout_sha256": "a" * 64, "stderr_sha256": "b" * 64}
    events = assemble_investigation_timeline(
        audit, [_finding("f1", "2026-07-15T11:00:00.000Z")]
    )
    rows = _rows(timeline_to_hoja(events))
    assert rows[0] == list(HOJA_HEADER)
    assert len(rows) == 3  # cabecera + tool_run + finding

    by_kind = {r[_IDX["Tipo de evento"]]: r for r in rows[1:]}

    tool = by_kind["Ejecución de herramienta"]
    assert tool[_IDX["N"]] == "1"
    assert tool[_IDX["Herramienta"]] == "tsk_fls"
    # El argv literal auditado, en la ÚLTIMA columna: es el dato más ancho y el
    # que antes empujaba fuera de pantalla a las columnas que se leen.
    assert tool[_IDX["Comando ejecutado (argv literal auditado)"]] == "fls -m / img.raw"
    assert tool[-1] == "fls -m / img.raw"
    assert tool[_IDX["Código de salida"]] == "0"
    assert tool[_IDX["Estado"]] == "Finalizada"
    assert tool[_IDX["Identificador del evento"]] == "r1"
    assert tool[_IDX["Marca temporal (UTC)"]].endswith("Z")  # UTC explícito
    # Lo que no aplica a una ejecución se DICE, no se deja en blanco.
    assert tool[_IDX["Severidad"]] == NO_APLICA
    assert tool[_IDX["Técnicas ATT&CK propuestas"]] == NO_APLICA

    finding = by_kind["Hallazgo"]
    assert finding[_IDX["N"]] == "2"
    assert finding[_IDX["Hallazgo"]] == "Persistencia detectada"
    assert finding[_IDX["Severidad"]] == "Alta"
    assert finding[_IDX["Detalle del hallazgo"]] == "run key sospechosa"
    assert finding[_IDX["Técnicas ATT&CK propuestas"]] == "T1547.001"
    assert finding[_IDX["Herramienta"]] == "RegRipper"
    assert finding[_IDX["Identificador del evento"]] == "f1"
    assert finding[_IDX["Comando ejecutado (argv literal auditado)"]] == NO_APLICA


def test_the_row_number_is_a_number_so_the_sheet_orders_by_value() -> None:
    """Con la numeración en texto, ordenar por la columna N pone la fila 10 delante
    de la 9. Entra como número para que ordene como se lee."""
    audit: list[dict[str, Any]] = []
    for n in range(12):
        audit.append(
            {
                "action": "tool_run_start",
                "ts_utc": f"2026-07-15T10:{n:02d}:00+00:00",
                "run_id": f"r{n}",
                "tool_id": "tsk_fls",
                "argv": ["fls", "img.raw"],
            }
        )
    events = assemble_investigation_timeline(audit, [])
    tabla = _hoja(timeline_to_hoja(events))[_PROCEDENCIA_FILAS + 2 :]
    assert [f[0] for f in tabla] == list(range(1, 13))


def test_running_tool_run_has_empty_exit_but_is_kept() -> None:
    audit = [
        {
            "action": "tool_run_start",
            "ts_utc": "2026-07-15T10:00:00+00:00",
            "run_id": "r1",
            "tool_id": "tsk_mmls",
            "argv": ["mmls", "img.raw"],
        },
    ]
    events = assemble_investigation_timeline(audit, [])
    row = _rows(timeline_to_hoja(events))[1]
    assert row[_IDX["Código de salida"]] == ""  # sin exit todavía
    assert row[_IDX["Estado"]] == "En curso"


def test_unparseable_timestamp_becomes_empty_but_row_is_not_dropped() -> None:
    audit = [
        {
            "action": "tool_run_start",
            "ts_utc": "not-a-date",
            "run_id": "r-bad",
            "tool_id": "tsk_fls",
            "argv": [],
        },
    ]
    events = assemble_investigation_timeline(audit, [])
    rows = _rows(timeline_to_hoja(events))
    assert len(rows) == 2  # cabecera + la fila, nunca descartada
    assert rows[1][_IDX["Marca temporal (UTC)"]] == ""


def test_an_unknown_vocabulary_value_travels_verbatim() -> None:
    """Un estado que no esté en la tabla de etiquetas NO se traduce a lo que se le
    parezca: es un dato del caso (RULE 2)."""
    audit = [
        {
            "action": "tool_run_start",
            "ts_utc": "2026-07-15T10:00:00+00:00",
            "run_id": "r1",
            "tool_id": "tsk_fls",
            "argv": ["fls", "img.raw"],
        },
        {
            "action": "tool_run_finish",
            "ts_utc": "2026-07-15T10:00:01+00:00",
            "run_id": "r1",
            "status": "cancelado_por_el_operador",
            "exit_code": 137,
        },
    ]
    events = assemble_investigation_timeline(audit, [])
    row = _rows(timeline_to_hoja(events))[1]
    assert row[_IDX["Estado"]] == "cancelado_por_el_operador"
    assert row[_IDX["Código de salida"]] == "137"
