"""Exportaciones de la cobertura ATT&CK: hoja de cálculo y ATT&CK Navigator layer.

Los gates que importan:
- La hoja se DERIVA de la cobertura real (propuestas del agente + dictámenes del
  perito), con nombres/tácticas del catálogo Enterprise, nunca inventados.
- Un caso con **cero** técnicas evaluadas da una hoja con su procedencia y la
  cabecera, 0 filas, nunca un error; y un layer válido sin celdas.
- La hoja se ABRE bien donde el perito la abre: un `.xlsx` real, con los acentos
  intactos y las celdas ya separadas, sin depender del idioma del sistema. El CSV
  anterior dependía de las dos cosas y perdía una: con la declaración ``sep=;``
  Excel deja de aplicar el BOM y «Exfiltración» sale «ExfiltraciÃ³n».
- El layer del Navigator es válido: formato 4.5, dominio enterprise-attack, colores
  por eje (dictamen del perito vs propuesta del agente), que no se funden.
"""

from __future__ import annotations

import io
import json
from typing import Any

import pytest
from openpyxl import load_workbook

from forensia.i18n import CATALOGO, t
from forensia.cases import CaseManager
from forensia.findings.store import FindingStore
from forensia.mitre.coverage import CoverageStore
from forensia.mitre.export import (
    NAVIGATOR_LAYER_VERSION,
    sin_dictamen,
    coverage_to_hoja,
    coverage_to_navigator_layer,
    hoja_header,
)

def _col(nombre_es: str) -> str:
    """El rótulo de una columna (o de una fila de procedencia) EN EL IDIOMA EN
    CURSO, nombrándola por su texto castellano.

    El test sigue leyéndose en castellano, que es la lengua del proyecto, pero no
    se rompe cuando la hoja se exporta en inglés: lo que fija es QUÉ columna, no
    con qué palabra se escribe.
    """
    for clave, entrada in CATALOGO.items():
        if entrada.get("es") == nombre_es:
            return t(clave)
    return nombre_es


#: Procedencia válida (UUID4) para hallazgos afirmativos — el store la exige
#: (anti-alucinación, RULE 2); estos tests ejercitan la export MITRE, no ese gate.
_RUN_ID = "11111111-1111-4111-8111-111111111111"


@pytest.fixture
def tmp_case(tmp_path) -> tuple[CaseManager, str]:
    cases = CaseManager(root=tmp_path / "cases")
    case = cases.create(name="Caso export", examiner="ramos", os_profile="windows")
    return cases, case.id


#: El bloque de procedencia de la hoja de cobertura (lo que hay antes de la tabla).
_PROCEDENCIA_FILAS = 8


def _sheet(blob: bytes) -> list[list[Any]]:
    """Las celdas tal cual las lee una hoja de cálculo, conservando el tipo."""
    ws = load_workbook(io.BytesIO(blob)).active
    return [[c.value for c in fila] for fila in ws.iter_rows()]


def _rows(blob: bytes) -> list[list[str]]:
    """Sólo la TABLA: la cabecera y sus filas, ya sin el preámbulo ni su línea
    vacía, en texto para comparar contra el contrato de la cabecera."""
    tabla = _sheet(blob)[_PROCEDENCIA_FILAS + 1 :]
    return [
        ["" if v is None else str(v) for v in fila]
        for fila in tabla
        if any(v is not None for v in fila)
    ]


def test_sheet_opens_as_a_spreadsheet_and_declares_its_provenance(tmp_case) -> None:
    """Un `.xlsx` real con su procedencia: ni codificación ni separador que
    negociar con el programa que la abre, que es lo que rompía el CSV."""
    cases, case_id = tmp_case
    coverage = CoverageStore(cases, FindingStore(cases))
    coverage.adjudicate(case_id, "T1055", "confirmada", "Motivo con acentuación.")
    blob = coverage_to_hoja(
        coverage.coverage(case_id),
        case_id=case_id,
        case_name="Caso export",
        exported_at="2026-08-06T13:05:42Z",
    )
    # Un `.xlsx` es un paquete OOXML: empieza por la firma de un zip.
    assert blob[:2] == b"PK"
    hoja = _sheet(blob)
    procedencia = dict((r[0], r[1]) for r in hoja[:_PROCEDENCIA_FILAS])
    assert procedencia[_col("Caso")] == "Caso export"
    assert procedencia[_col("Identificador del caso")] == case_id
    assert procedencia[_col("Exportado (UTC)")] == "2026-08-06T13:05:42Z"
    assert procedencia[_col("Técnicas en la hoja")] == "1"
    # El texto viaja en UTF-8 dentro del paquete, no transliterado.
    assert "ó" in _rows(blob)[1][-1]
    # Una línea vacía separa la procedencia de la tabla: es el contrato para
    # quien lea la hoja con un programa.
    assert all(v is None for v in hoja[_PROCEDENCIA_FILAS])
    assert _rows(blob)[0] == list(hoja_header())


def test_sheet_empty_case_is_header_only(tmp_case) -> None:
    """Cero técnicas evaluadas → cabecera y 0 filas, no un error."""
    cases, case_id = tmp_case
    coverage = CoverageStore(cases, FindingStore(cases))
    rows = _rows(coverage_to_hoja(coverage.coverage(case_id)))
    assert rows == [list(hoja_header())]


def test_sheet_row_for_an_agent_proposal(tmp_case) -> None:
    cases, case_id = tmp_case
    findings = FindingStore(cases)
    coverage = CoverageStore(cases, findings)
    f = findings.append(case_id, {
        "title": "Inyección",
        "summary": "malfind: región RWX.",
        "severity": "high",
        "run_id": _RUN_ID,
        "mitre_hints": ["T1055"],
    })
    rows = _rows(coverage_to_hoja(coverage.coverage(case_id)))
    assert rows[0] == list(hoja_header())
    idx = {name: i for i, name in enumerate(hoja_header())}
    body = {r[idx[_col("ID de la técnica")]]: r for r in rows[1:]}
    row = body["T1055"]
    assert row[idx[_col('N')]] == "1"
    # Nombre y táctica salen del catálogo Enterprise, NO inventados.
    assert row[idx[_col('Técnica')]] == "Process Injection"
    assert row[idx[_col('ID de la táctica')]] == "TA0005"
    assert row[idx[_col('Táctica')]] != ""
    assert row[idx[_col('Propuesta por el análisis')]] == t("sheet.yes")
    assert row[idx[_col('Hallazgos que la proponen')]] == "1"
    # Sin dictamen se DICE que está pendiente: en blanco se leería «no aplica».
    assert row[idx[_col('Veredicto del perito')]] == sin_dictamen()
    assert row[idx[_col('Identificadores de hallazgo')]] == f.id
    assert row[idx[_col('Celda de la matriz')]] == "T1055"


def test_sheet_carries_the_examiner_verdict_and_rationale(tmp_case) -> None:
    cases, case_id = tmp_case
    findings = FindingStore(cases)
    coverage = CoverageStore(cases, findings)
    f = findings.append(case_id, {
        "title": "Inyección", "summary": "malfind.", "severity": "high",
        "run_id": _RUN_ID,
        "mitre_hints": ["T1055"],
    })
    coverage.adjudicate(
        case_id, "T1055", "confirmada", "RWX + shellcode.", related_finding_ids=[f.id]
    )
    rows = _rows(coverage_to_hoja(coverage.coverage(case_id)))
    idx = {name: i for i, name in enumerate(hoja_header())}
    row = {r[idx[_col("ID de la técnica")]]: r for r in rows[1:]}["T1055"]
    assert row[idx[_col('Veredicto del perito')]] == "Confirmada"
    assert row[idx[_col('Motivo del veredicto')]] == "RWX + shellcode."
    assert row[idx[_col('Propuesta por el análisis')]] == t("sheet.yes")
    assert row[idx[_col('Fecha del veredicto (UTC)')]] != ""


def test_sheet_row_for_a_verdict_without_an_agent_proposal(tmp_case) -> None:
    """El perito puede dictaminar una técnica que el agente no propuso."""
    cases, case_id = tmp_case
    coverage = CoverageStore(cases, FindingStore(cases))
    coverage.adjudicate(case_id, "T1070", "descartada", "Los logs están intactos.")
    rows = _rows(coverage_to_hoja(coverage.coverage(case_id)))
    idx = {name: i for i, name in enumerate(hoja_header())}
    row = {r[idx[_col("ID de la técnica")]]: r for r in rows[1:]}["T1070"]
    assert row[idx[_col('Propuesta por el análisis')]] == "No"
    assert row[idx[_col('Veredicto del perito')]] == "Descartada"
    assert row[idx[_col('Identificadores de hallazgo')]] == ""


def test_punctuation_in_a_rationale_survives_without_escaping(tmp_case) -> None:
    """El motivo de un dictamen es prosa del perito y puede llevar el separador y
    comillas. En una hoja no hay nada que escapar: la celda es una celda."""
    cases, case_id = tmp_case
    coverage = CoverageStore(cases, FindingStore(cases))
    coverage.adjudicate(
        case_id, "T1055", "sospechosa", 'Motivo; con punto y coma y "comillas".'
    )
    rows = _rows(coverage_to_hoja(coverage.coverage(case_id)))
    idx = {name: i for i, name in enumerate(hoja_header())}
    row = {r[idx[_col("ID de la técnica")]]: r for r in rows[1:]}["T1055"]
    assert row[idx[_col('Motivo del veredicto')]] == 'Motivo; con punto y coma y "comillas".'


# ── ATT&CK Navigator layer ───────────────────────────────────────────────────


def test_navigator_layer_is_valid_and_empty_when_no_coverage(tmp_case) -> None:
    cases, case_id = tmp_case
    coverage = CoverageStore(cases, FindingStore(cases))
    layer = coverage_to_navigator_layer(
        coverage.coverage(case_id), case_id=case_id, case_name="Caso export"
    )
    json.dumps(layer)  # serializable
    assert layer["versions"]["layer"] == NAVIGATOR_LAYER_VERSION == "4.5"
    assert layer["domain"] == "enterprise-attack"
    assert layer["techniques"] == []
    assert "Caso export" in layer["name"]


def test_navigator_layer_colors_by_axis(tmp_case) -> None:
    cases, case_id = tmp_case
    findings = FindingStore(cases)
    coverage = CoverageStore(cases, findings)
    # Propuesta del agente sin dictaminar → color de propuesta.
    findings.append(case_id, {
        "title": "Persistencia", "summary": "run key.", "severity": "medium",
        "run_id": _RUN_ID,
        "mitre_hints": ["T1547.001"],
    })
    # Dictamen del perito → color de veredicto.
    coverage.adjudicate(case_id, "T1055", "confirmada", "RWX + shellcode.")

    layer = coverage_to_navigator_layer(
        coverage.coverage(case_id), case_id=case_id, case_name="c"
    )
    by_id = {t["techniqueID"]: t for t in layer["techniques"]}
    # La sub-técnica de la semilla se pinta en su celda Enterprise padre.
    assert "T1547" in by_id
    assert by_id["T1547"]["color"] == "#2f6fed"  # propuesta (azul)
    assert by_id["T1055"]["color"] == "#c1121f"  # confirmada (rojo)
    assert "Dictamen del perito: Confirmada" in by_id["T1055"]["comment"]
    assert all(t["enabled"] for t in layer["techniques"])


def test_navigator_layer_verdict_wins_over_proposal_color(tmp_case) -> None:
    """Si una técnica está propuesta Y dictaminada, el color lo manda el dictamen."""
    cases, case_id = tmp_case
    findings = FindingStore(cases)
    coverage = CoverageStore(cases, findings)
    findings.append(case_id, {
        "title": "Inyección", "summary": "malfind.", "severity": "high",
        "run_id": _RUN_ID,
        "mitre_hints": ["T1055"],
    })
    coverage.adjudicate(case_id, "T1055", "sospechosa", "Indicio.")
    layer = coverage_to_navigator_layer(
        coverage.coverage(case_id), case_id=case_id, case_name="c"
    )
    t = {t["techniqueID"]: t for t in layer["techniques"]}["T1055"]
    assert t["color"] == "#e08a00"  # sospechosa (ámbar), no el azul de propuesta
    assert "Propuesta del agente" in t["comment"]  # pero la propuesta también consta
