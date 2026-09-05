"""El envoltorio de un export: `agentopsy.export_hoja`.

Aquí se fija lo que hace que la hoja se pueda ABRIR y ADJUNTAR, que es distinto de
lo que hace que sus datos sean correctos (eso lo fijan `test_mitre_export` y
`test_timeline_export`, cada uno sobre su dominio):

- que el fichero sea un `.xlsx` de verdad, no un texto con otra extensión;
- que una celda que empiece por «=» NO se convierta en fórmula, porque ese texto
  puede venir de la evidencia, que es hostil (SECURITY INVARIANTS);
- que la hoja venga presentada como un anexo y no como un volcado: cabecera
  fijada, filtros, anchos acotados y ajuste de impresión.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from agentopsy.export_hoja import (
    EXTENSION,
    MEDIA_TYPE,
    _ANCHO_MAX,
    build_workbook,
    export_basename,
    unir,
)

_PROCEDENCIA = [
    ("Agentopsy", "Hoja de prueba"),
    ("Caso", "Caso Ñandú"),
]
_CABECERA = ["N", "Texto", "Número"]


def _ws(blob: bytes):
    return load_workbook(io.BytesIO(blob)).active


def test_the_file_is_a_real_spreadsheet_package() -> None:
    blob = build_workbook(
        procedencia=_PROCEDENCIA, cabecera=_CABECERA, filas=[[1, "hola", 7]],
        titulo="Prueba",
    )
    assert blob[:2] == b"PK"  # paquete OOXML (zip), no texto con otra extensión
    assert EXTENSION == "xlsx"
    assert MEDIA_TYPE.endswith("spreadsheetml.sheet")
    assert _ws(blob).title == "Prueba"


def test_hostile_evidence_never_becomes_a_formula() -> None:
    """Un `summary` o un `argv` salen de la imagen del sospechoso y pueden empezar
    por «=». Una hoja de cálculo interpretaría eso como fórmula: es la inyección de
    fórmulas clásica, y aquí el texto que la dispara viene literalmente del disco
    analizado. La celda tiene que enseñar el byte que había, no ejecutarlo."""
    hostil = [
        '=cmd|" /c calc"!A1',
        "=1+1",
        "+HYPERLINK(\"http://evil\")",
        "@SUM(A1:A9)",
        "-2+3",
    ]
    blob = build_workbook(
        procedencia=_PROCEDENCIA,
        cabecera=["N", "Texto"],
        filas=[[i, t] for i, t in enumerate(hostil, start=1)],
        titulo="Prueba",
    )
    ws = _ws(blob)
    primera = len(_PROCEDENCIA) + 3  # procedencia + línea vacía + cabecera
    for i, esperado in enumerate(hostil):
        celda = ws.cell(row=primera + i, column=2)
        assert celda.data_type == "s", f"{esperado!r} entró como fórmula"
        assert celda.value == esperado  # y el texto llega entero, sin recortar


def test_numbers_stay_numbers_and_everything_else_stays_text() -> None:
    """Un entero entra como número para que la columna ordene por valor; un
    identificador entra como texto para que no se redondee ni se convierta a
    notación científica."""
    blob = build_workbook(
        procedencia=_PROCEDENCIA,
        cabecera=_CABECERA,
        filas=[[1, "0001234500067890", 42], [2, "T1055", 0]],
        titulo="Prueba",
    )
    ws = _ws(blob)
    primera = len(_PROCEDENCIA) + 3
    assert ws.cell(row=primera, column=1).value == 1
    assert ws.cell(row=primera, column=3).value == 42
    assert ws.cell(row=primera, column=2).value == "0001234500067890"
    assert ws.cell(row=primera + 1, column=3).value == 0  # un cero real, no vacío


def test_the_sheet_is_presented_as_an_annex() -> None:
    """Cabecera a la vista al desplazarse, filtros para ordenar sin tocar el
    fichero, anchos acotados y ajuste de impresión: es lo que separa un anexo de
    un volcado."""
    filas = [[n, "x" * 400, n] for n in range(1, 6)]
    blob = build_workbook(
        procedencia=_PROCEDENCIA, cabecera=_CABECERA, filas=filas, titulo="Prueba",
    )
    ws = _ws(blob)
    fila_cabecera = len(_PROCEDENCIA) + 2
    assert ws.freeze_panes == f"A{fila_cabecera + 1}"
    assert ws.auto_filter.ref == f"A{fila_cabecera}:C{fila_cabecera + len(filas)}"
    # Una celda larguísima (un argv auditado pasa de cien caracteres) no puede
    # llevarse el ancho de la hoja por delante; el contenido sigue entero.
    assert ws.column_dimensions["B"].width <= _ANCHO_MAX
    assert ws.print_titles == f"'Prueba'!${fila_cabecera}:${fila_cabecera}"
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True


def test_an_empty_table_carries_no_filter() -> None:
    """Con 0 filas, Excel avisa de un rango de filtro sin datos y la hoja se
    abriría con una advertencia. La cabecera sí se queda fijada."""
    blob = build_workbook(
        procedencia=_PROCEDENCIA, cabecera=_CABECERA, filas=[], titulo="Prueba",
    )
    ws = _ws(blob)
    assert ws.auto_filter.ref is None
    assert ws.freeze_panes == f"A{len(_PROCEDENCIA) + 3}"


def test_the_sheet_name_survives_what_excel_forbids() -> None:
    """Excel rechaza `[]:*?/\\` en el nombre de una pestaña y la corta a 31."""
    ws = _ws(
        build_workbook(
            procedencia=_PROCEDENCIA,
            cabecera=_CABECERA,
            filas=[],
            titulo="Cobertura [ATT&CK]: 2026/2027 " + "x" * 40,
        )
    )
    assert len(ws.title) <= 31
    assert not set(ws.title) & set("[]:*?/\\")


def test_export_basename_names_the_tool_the_kind_the_case_and_the_moment() -> None:
    nombre = export_basename("Caso Ñandú", "timeline", exported_at="2026-08-06T13:05:42Z")
    assert nombre == "agentopsy-timeline-caso-nandu-20260806T130542Z"
    assert nombre.isascii()  # viaja en Content-Disposition


def test_export_basename_of_a_name_that_transliterates_to_nothing() -> None:
    """Un nombre escrito íntegro en un alfabeto no latino deja la herramienta, el
    tipo y la fecha, nunca un nombre inventado (RULE 2)."""
    nombre = export_basename("日本語", "timeline", exported_at="2026-08-06T13:05:42Z")
    assert nombre == "agentopsy-timeline-20260806T130542Z"


def test_unir_leaves_out_what_is_absent() -> None:
    assert unir(["T1055", None, "", "T1059"]) == "T1055 | T1059"
