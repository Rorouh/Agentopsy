"""Los dos endpoints de exportación, por HTTP (2026-08-06).

Las funciones que construyen las hojas se prueban en ``test_mitre_export.py`` y
``test_timeline_export.py``. Aquí se prueba lo que solo se ve al atravesar la
superficie, que es donde una hoja bien construida todavía puede llegar
inservible:

- el ``Content-Type`` y el ``Content-Disposition``: el nombre lleva la
  herramienta, el tipo, el caso y la marca temporal, y **está transcrito a
  ASCII**, porque una cabecera HTTP no admite otra cosa y un caso llamado
  «Análisis Ñandú» reventaría la respuesta entera;
- que los bytes que salen por el cable son un `.xlsx` de verdad y que el
  ``Content-Type`` lo declara, que es lo que hace que el navegador y el sistema lo
  ofrezcan a la hoja de cálculo en lugar de a un editor de texto;
- que un caso inexistente da 404 y no una hoja vacía (RULE 2).
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from forensia.i18n import t
from forensia.cases.manager import CaseManager
from forensia.export_hoja import MEDIA_TYPE
from forensia.findings.store import FindingStore
from forensia.mitre.coverage import CoverageStore
from forensia.server import create_app

PORT = 51033

#: Procedencia válida (UUID4) para un hallazgo afirmativo: el store la exige.
_RUN_ID = "22222222-2222-4222-8222-222222222222"

#: Un nombre con acentos, eñe y un carácter que no sobrevive a la transcripción.
#: Es el caso que rompería la cabecera de descarga si el nombre viajara tal cual.
_NOMBRE = "Análisis Ñandú 2026 · exfiltración"


@pytest.fixture
def cases(tmp_path, monkeypatch) -> CaseManager:
    """Un `CaseManager` aislado, inyectado en los routers que exportan."""
    store = CaseManager(root=tmp_path / "cases")

    import forensia.mitre.coverage as coverage_mod
    import forensia.routers.mitre as mitre_router
    import forensia.routers.timeline as timeline_router
    import forensia.timeline.builder as builder

    findings = FindingStore(store)
    coverage = CoverageStore(store, findings)
    monkeypatch.setattr(mitre_router, "case_manager", store)
    monkeypatch.setattr(mitre_router, "coverage_store", coverage)
    monkeypatch.setattr(coverage_mod, "coverage_store", coverage)
    monkeypatch.setattr(timeline_router, "case_manager", store)
    # El builder del timeline resuelve el caso por su PROPIO singleton, no por el
    # del router: en producción es el mismo objeto, en el test hay que decírselo.
    monkeypatch.setattr(builder, "case_manager", store)
    monkeypatch.setattr(builder, "finding_store", findings)
    return store


@pytest.fixture
def client(cases) -> TestClient:
    return TestClient(create_app(PORT), base_url=f"http://127.0.0.1:{PORT}")


@pytest.fixture
def auth(client) -> dict[str, str]:
    return {"X-Forensia-Token": client.app.state.token}


def _filename(response) -> str:
    disposition = response.headers["Content-Disposition"]
    assert disposition.startswith('attachment; filename="')
    return disposition.split('"')[1]


def _texto(response) -> str:
    """Todo el texto de la hoja descargada, para buscar dentro sin saber la
    posición exacta de cada celda."""
    ws = load_workbook(io.BytesIO(response.content)).active
    return "\n".join(
        "\t".join("" if c.value is None else str(c.value) for c in fila)
        for fila in ws.iter_rows()
    )


@pytest.mark.parametrize(
    ("ruta", "kind"),
    [("mitre/export.xlsx", "mitre-attack"), ("timeline/export.xlsx", "timeline")],
)
def test_the_downloaded_sheet_is_named_in_ascii_and_is_a_real_workbook(
    client, auth, cases, ruta, kind
) -> None:
    case = cases.create(name=_NOMBRE, examiner="perito", os_profile="windows")

    res = client.get(f"/api/cases/{case.id}/{ruta}", headers=auth)

    assert res.status_code == 200, res.text
    assert res.headers["Content-Type"] == MEDIA_TYPE
    nombre = _filename(res)
    assert nombre.isascii(), "una cabecera HTTP no admite otra cosa"
    assert nombre.startswith(f"agentopsy-{kind}-analisis-nandu-2026-exfiltracion-")
    assert nombre.endswith(".xlsx")
    # Lo que sale por el cable se abre como libro, no es texto con otra extensión.
    assert res.content[:2] == b"PK"
    # El nombre del caso llega íntegro DENTRO de la hoja: lo que se transcribe a
    # ASCII es el nombre del fichero, nunca el dato.
    assert _NOMBRE in _texto(res)


def test_the_navigator_layer_is_named_the_same_way(client, auth, cases) -> None:
    case = cases.create(name=_NOMBRE, examiner="perito", os_profile="windows")

    res = client.get(f"/api/cases/{case.id}/mitre/navigator", headers=auth)

    assert res.status_code == 200, res.text
    nombre = _filename(res)
    assert nombre.isascii()
    assert nombre.startswith("agentopsy-mitre-navigator-analisis-nandu-2026-exfiltracion-")
    assert nombre.endswith(".json")
    assert _NOMBRE in res.json()["name"]


def test_the_sheet_carries_the_real_coverage_of_the_case(client, auth, cases) -> None:
    """No es una hoja de adorno: la fila sale de la cobertura persistida."""
    case = cases.create(name="Caso con cobertura", examiner="perito", os_profile="windows")
    findings = FindingStore(cases)
    coverage = CoverageStore(cases, findings)
    findings.append(case.id, {
        "title": "Inyección",
        "summary": "malfind: región RWX.",
        "severity": "high",
        "run_id": _RUN_ID,
        "mitre_hints": ["T1055"],
    })
    assert coverage.coverage(case.id), "precondición: la cobertura existe"

    texto = _texto(client.get(f"/api/cases/{case.id}/mitre/export.xlsx", headers=auth))

    assert "T1055" in texto
    assert "Process Injection" in texto
    assert f'{t("mitreSheet.prov.techniques")}\t1' in texto


@pytest.mark.parametrize(
    "ruta", ["mitre/export.xlsx", "timeline/export.xlsx", "mitre/navigator"]
)
def test_a_case_that_does_not_exist_is_a_404_not_an_empty_sheet(
    client, auth, ruta
) -> None:
    res = client.get(f"/api/cases/00000000-0000-4000-8000-000000000000/{ruta}", headers=auth)
    assert res.status_code == 404, res.text
